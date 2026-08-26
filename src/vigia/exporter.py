from __future__ import annotations

import csv
import os
from datetime import datetime
from typing import Any

from vigia.store import Store

COLUNAS = [
    "data",
    "perfil",
    "objeto",
    "regra",
    "valor_estimado",
    "controle_pncp",
]


def _historico_para_rows(
    store: Store,
    perfil: str | None = None,
    desde: str | None = None,
    ate: str | None = None,
) -> list[dict[str, Any]]:
    historico = store.historico_alertas(perfil=perfil, limite=10000)

    filtrado = []
    for item in historico:
        data_str = item.get("visto_em", "")[:10]
        if desde and data_str < desde:
            continue
        if ate and data_str > ate:
            continue
        filtrado.append({
            "data": item.get("visto_em", ""),
            "perfil": item.get("perfil", ""),
            "objeto": (item.get("objeto") or "").replace("\n", " ").strip(),
            "regra": item.get("regra", ""),
            "valor_estimado": item.get("valor") or "",
            "controle_pncp": item.get("controle", ""),
        })
    return filtrado


def _gerar_caminho(
    sufixo: str, extensao: str, caminho_saida: str | None = None
) -> str:
    if caminho_saida:
        return caminho_saida
    agora = datetime.now().strftime("%Y-%m-%d_%H%M")
    return f"vigia_export{sufixo}_{agora}.{extensao}"


def exportar_csv(
    store: Store,
    perfil: str | None = None,
    desde: str | None = None,
    ate: str | None = None,
    caminho_saida: str | None = None,
) -> str:
    """Exporta alertas para CSV. Retorna caminho do arquivo."""
    rows = _historico_para_rows(store, perfil, desde, ate)
    sufixo = f"_{perfil}" if perfil else ""
    caminho = _gerar_caminho(sufixo, "csv", caminho_saida)

    os.makedirs(os.path.dirname(caminho) or ".", exist_ok=True)

    with open(caminho, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUNAS)
        writer.writeheader()
        writer.writerows(rows)

    return caminho


def exportar_xlsx(
    store: Store,
    perfil: str | None = None,
    desde: str | None = None,
    ate: str | None = None,
    caminho_saida: str | None = None,
) -> str:
    """Exporta alertas para XLSX. Retorna caminho do arquivo."""
    rows = _historico_para_rows(store, perfil, desde, ate)
    sufixo = f"_{perfil}" if perfil else ""
    caminho = _gerar_caminho(sufixo, "xlsx", caminho_saida)

    try:
        from openpyxl import Workbook  # type: ignore[import-untyped]
        from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
    except ImportError as exc:
        raise ImportError(
            "Para exportar XLSX, instale openpyxl: uv add openpyxl"
        ) from exc

    os.makedirs(os.path.dirname(caminho) or ".", exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Alertas"

    header_fill = PatternFill(
        start_color="1e293b", end_color="1e293b", fill_type="solid"
    )
    header_font = Font(color="38bdf8", bold=True)
    for col, nome in enumerate(COLUNAS, 1):
        cell = ws.cell(
            row=1, column=col, value=nome.replace("_", " ").title()
        )
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, row in enumerate(rows, 2):
        for col, key in enumerate(COLUNAS, 1):
            ws.cell(row=row_idx, column=col, value=str(row.get(key, "")))

    for col in range(1, len(COLUNAS) + 1):
        max_len = max(
            len(str(ws.cell(row=r, column=col).value or ""))
            for r in range(1, min(len(rows) + 2, 100))
        )
        ws.column_dimensions[
            ws.cell(row=1, column=col).column_letter
        ].width = min(max_len + 2, 60)

    wb.save(caminho)
    return caminho
